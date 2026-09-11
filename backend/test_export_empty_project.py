"""TC 가 하나도 없어도 시트 분리 내보내기가 파일을 만든다.

시트 분리 모드는 기본 시트를 지우고 TC 의 sheet_name 별로 탭을 만든다. TC 가 0건이면
탭이 하나도 남지 않아 openpyxl 이 저장 단계에서 IndexError("At least one sheet must be
visible") 를 던지고, 사용자에게는 500 으로 보인다. 시트만 만들어 두고 TC 를 아직 안 넣은
신규 프로젝트에서 그대로 물린다.

통합 모드는 시트를 항상 한 장 쓰므로 같은 상황에서 헤더만 있는 파일이 나온다. 두 모드의
결과가 갈리지 않아야 한다.

실행: cd backend && python -m pytest test_export_empty_project.py -v
"""
import asyncio
import io
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


class _Project:
    """export 가 쓰는 것은 name 뿐이다."""

    id = 1
    name = "빈프로젝트"


class _TC:
    """TestCase 에서 export 가 읽는 칸만 갖춘 더미."""

    def __init__(self, no, tc_id, sheet_name):
        self.no = no
        self.tc_id = tc_id
        self.sheet_name = sheet_name
        # models.TestCase 의 실제 칸만 둔다. export 가 읽는 것은 이 중 일부다.
        for attr in (
            "type", "category", "depth1", "depth2", "priority", "test_type",
            "precondition", "test_steps", "expected_result",
            "r1", "r2", "r3", "remarks",
        ):
            setattr(self, attr, "")


def _body(response):
    """StreamingResponse 본문을 bytes 로 모은다. body_iterator 는 async 라 루프를 돌린다."""

    async def collect():
        chunks = []
        async for chunk in response.body_iterator:
            chunks.append(chunk if isinstance(chunk, bytes) else chunk.encode())
        return b"".join(chunks)

    return asyncio.run(collect())


@pytest.fixture
def export():
    from services.export_service import export_testcases_excel

    return export_testcases_excel


def test_split_sheets_with_no_testcases(export):
    """TC 0건 + 시트 분리: 저장이 터지지 않고 시트가 최소 한 장 남는다."""
    res = export(_Project(), [], True, False)
    wb = load_workbook(io.BytesIO(_body(res)), read_only=True)
    assert len(wb.sheetnames) >= 1, "시트가 한 장도 없으면 엑셀이 열리지 않는다"
    wb.close()


def test_merged_sheet_with_no_testcases(export):
    """TC 0건 + 통합: 종전대로 헤더만 있는 파일이 나온다."""
    res = export(_Project(), [], False, False)
    wb = load_workbook(io.BytesIO(_body(res)), read_only=True)
    assert wb.sheetnames == ["Test Cases"]
    wb.close()


def test_split_sheets_keeps_tabs_per_sheet_name(export):
    """TC 가 있으면 sheet_name 마다 탭이 생긴다. 엑셀 금지 문자는 바뀐다."""
    tcs = [
        _TC(1, "AUTH-001", "인증/보안"),
        _TC(2, "AUTH-002", "인증/보안"),
        _TC(3, "HOME-001", "홈"),
    ]
    res = export(_Project(), tcs, True, False)
    wb = load_workbook(io.BytesIO(_body(res)), read_only=True)
    assert wb.sheetnames == ["인증-보안", "홈"]
    wb.close()

def test_split_sheets_titles_obey_excel_limits(export):
    """시트명은 31자 이하이고 서로 달라야 한다.

    31자까지 같은 이름 둘은 자른 뒤 충돌한다. openpyxl 은 중복을 만나면 뒤에 번호를 붙이는데,
    그 번호가 31자 상한 밖으로 나가면 엑셀이 파일을 복구 대상으로 본다. 실측에서 32자짜리
    탭이 나왔다.
    """
    long1 = "가" * 30 + "끝1"
    long2 = "가" * 30 + "끝2"
    tcs = [_TC(1, "E-001", long1), _TC(2, "E-002", long2)]
    res = export(_Project(), tcs, True, False)
    wb = load_workbook(io.BytesIO(_body(res)), read_only=True)
    names = wb.sheetnames
    wb.close()
    assert len(names) == 2, names
    assert len(set(names)) == 2, f"시트명이 겹치면 안 된다: {names}"
    assert all(len(n) <= 31 for n in names), f"31자를 넘는 탭: {[(n, len(n)) for n in names]}"


def test_split_sheets_title_with_only_banned_chars(export):
    """금지 문자만 남는 이름도 탭 하나로 살아남는다."""
    tcs = [_TC(1, "E-003", "*?[]:")]
    res = export(_Project(), tcs, True, False)
    wb = load_workbook(io.BytesIO(_body(res)), read_only=True)
    names = wb.sheetnames
    wb.close()
    assert len(names) == 1, names
    assert names[0].strip(), f"빈 이름이면 안 된다: {names}"
