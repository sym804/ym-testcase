"""런 엑셀의 No 도 화면과 같은 순번이다.

화면은 지금 보고 있는 목록의 순번을 1 부터 매기는데 내보내기만 저장된 no 를
그대로 썼다. 같은 수행인데 화면과 파일의 번호가 갈렸다.

실행: cd backend && TEST_PORT=8009 TEST_BASE_URL=http://127.0.0.1:8009 python -m pytest test_run_export_numbering.py -v
"""
import io
import os

import pytest
import requests
from openpyxl import load_workbook

BASE = os.getenv("TEST_BASE_URL", "http://127.0.0.1:8008")
ADMIN_PW = os.getenv("TEST_ADMIN_PASSWORD", "test1234")

if BASE.endswith(":8008") and os.getenv("ALLOW_DEV_DB") != "1":
    pytest.skip(
        "개발 서버(8008)의 실 DB 오염 방지를 위해 건너뜀. "
        "격리 실행: TEST_PORT=8009 TEST_BASE_URL=http://127.0.0.1:8009 pytest test_run_export_numbering.py",
        allow_module_level=True,
    )


def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": ADMIN_PW})
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


def _add_tc(token, pid, no, tc_id, sheet):
    r = requests.post(f"{BASE}/api/projects/{pid}/testcases", headers=auth(token), json={
        "no": no, "tc_id": tc_id, "category": sheet, "test_steps": "1. 실행",
        "expected_result": "성공", "sheet_name": sheet, "priority": "High",
    })
    assert r.status_code == 201, r.text
    return r.json()["id"]


@pytest.fixture
def project(token):
    """번호에 구멍이 있는 시트 둘. 시트 등록 순서는 결제 -> 로그인 이다."""
    h = auth(token)
    r = requests.post(f"{BASE}/api/projects", headers=h, json={"name": "__run_export_no__"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]

    for sheet in ("결제", "로그인"):
        rs = requests.post(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=h,
                           json={"name": sheet, "parent_id": None, "is_folder": False})
        assert rs.status_code in (200, 201), rs.text

    for no, tc in ((4, "TC-결제-A"), (9, "TC-결제-B")):
        _add_tc(token, pid, no, tc, "결제")
    for no, tc in ((2, "TC-로그인-A"), (7, "TC-로그인-B")):
        _add_tc(token, pid, no, tc, "로그인")

    yield pid
    requests.delete(f"{BASE}/api/projects/{pid}", headers=h)


def _export_rows(token, pid, run_id):
    r = requests.get(f"{BASE}/api/projects/{pid}/testruns/{run_id}/export", headers=auth(token))
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    return [(row[0], row[1]) for row in ws.iter_rows(min_row=2, values_only=True)]


def _make_run(token, pid, body):
    r = requests.post(f"{BASE}/api/projects/{pid}/testruns", headers=auth(token), json=body)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _report_rows(token, pid, run_id):
    r = requests.get(f"{BASE}/api/projects/{pid}/reports/excel",
                     headers=auth(token), params={"run_id": run_id})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Results"]
    return [(row[0], row[1]) for row in ws.iter_rows(min_row=2, values_only=True)]


def test_report_excel_matches_run_excel(token, project):
    """리포트 엑셀과 수행 엑셀이 같은 번호, 같은 순서를 낸다.

    같은 수행을 두 파일로 뽑을 수 있다. 한쪽만 고치면 같은 행이 다른 번호를 단다.
    """
    run_id = _make_run(token, project, {"name": "리포트 대조"})

    assert _report_rows(token, project, run_id) == _export_rows(token, project, run_id)


def test_export_numbers_from_one(token, project):
    """시트 하나로 범위를 잡은 수행의 No 는 1, 2 다."""
    run_id = _make_run(token, project, {"name": "결제 회귀", "sheet_names": ["결제"]})

    rows = _export_rows(token, project, run_id)

    assert [no for no, _ in rows] == [1, 2]
    assert [tc for _, tc in rows] == ["TC-결제-A", "TC-결제-B"]


def _leaf_tabs(token, pid):
    """화면 시트 탭에 놓이는 차례. 폴더를 걷어 낸 잎 시트만 깊이 우선으로 편다."""
    r = requests.get(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=auth(token))
    assert r.status_code == 200, r.text

    out = []

    def walk(nodes):
        for n in nodes:
            if not n["is_folder"] and (n.get("tc_count") or 0) > 0:
                out.append(n["name"])
            walk(n.get("children") or [])

    walk(r.json())
    return out


def test_export_order_matches_sheet_tabs(token, project):
    """내보내기 순서는 화면 시트 탭 순서와 같아야 한다.

    두 곳이 시트 순서를 따로 계산한다. 한쪽만 고치면 조용히 갈라지므로 여기서 묶는다.
    폴더 밑에 시트를 넣어도 어긋나지 않는지 함께 본다.
    """
    pid = project
    h = auth(token)
    rf = requests.post(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=h,
                       json={"name": "회귀", "parent_id": None, "is_folder": True})
    assert rf.status_code in (200, 201), rf.text
    folder_id = rf.json().get("id")
    rs = requests.post(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=h,
                       json={"name": "정산", "parent_id": folder_id, "is_folder": False})
    assert rs.status_code in (200, 201), rs.text
    _add_tc(token, pid, 1, "TC-정산-A", "정산")

    run_id = _make_run(token, pid, {"name": "전체"})
    rows = _export_rows(token, pid, run_id)

    tabs = _leaf_tabs(token, pid)
    assert [no for no, _ in rows] == list(range(1, len(rows) + 1))
    # TC ID 접두어가 그 TC 의 시트 이름이다. 시트가 바뀌는 자리만 뽑아 비교한다.
    seen = []
    for _, tc_id in rows:
        name = tc_id.split("-")[1]
        if not seen or seen[-1] != name:
            seen.append(name)
    assert seen == tabs


def test_export_follows_sheet_order(token, project):
    """여러 시트를 담은 수행은 시트 순서대로 세운 뒤 이어서 매긴다."""
    run_id = _make_run(token, project, {"name": "전체 회귀"})

    rows = _export_rows(token, project, run_id)

    assert [no for no, _ in rows] == [1, 2, 3, 4]
    # 시트 등록 순서가 결제 -> 로그인 이므로 결제가 먼저다. 저장된 no 순이 아니다.
    assert [tc for _, tc in rows] == [
        "TC-결제-A", "TC-결제-B", "TC-로그인-A", "TC-로그인-B",
    ]
