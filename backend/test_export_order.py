"""TC 목록 엑셀의 행 차례

`no` 는 v1.5.0.0 이후 **시트 안 순번**이라 시트마다 1 부터 다시 시작한다. 그래서
`order_by(no)` 만 걸면 시트가 한 줄씩 번갈아 나온다(A-1, B-1, A-2, B-2 ...).
시트가 9개면 No 칸에 1 이 아홉 번 찍힌다. 통합 모드가 화면 기본값이다.

수행 엑셀과 리포트 엑셀은 SYM-44 때 `leaf_sheet_order` 로 통일했는데 TC 목록
엑셀만 그 통일에서 빠졌다.

실행: cd backend && TEST_PORT=8099 TEST_BASE_URL=http://127.0.0.1:8099 python -m pytest test_export_order.py -v
"""
import io
import os

import pytest
import requests
from openpyxl import load_workbook

BASE = os.getenv("TEST_BASE_URL", "http://127.0.0.1:8008")
ADMIN_PW = os.getenv("TEST_ADMIN_PASSWORD", "test1234")

import dev_db_guard

if dev_db_guard.DEV_DB_AT_RISK:
    pytest.skip(dev_db_guard.SKIP_REASON, allow_module_level=True)


def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": ADMIN_PW})
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture
def project(token):
    """시트 셋(가, 나, 다)에 TC 2건씩. 시트 차례는 만든 순서다."""
    h = auth(token)
    r = requests.post(f"{BASE}/api/projects", headers=h, json={"name": "__export_order__"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]

    # ★생성 순서를 이름 순서와 어긋나게 둔다. 둘이 같으면 `order_by(no)` 로
    #   되돌려도 탭 차례가 우연히 맞아 시트 분리 테스트가 통과해 버린다.
    for sheet in ("다시트", "가시트", "나시트"):
        rs = requests.post(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=h,
                           json={"name": sheet, "parent_id": None, "is_folder": False})
        assert rs.status_code in (200, 201), rs.text

    for sheet in ("다시트", "가시트", "나시트"):
        for i in (1, 2):
            rr = requests.post(f"{BASE}/api/projects/{pid}/testcases", headers=h, json={
                "no": i, "tc_id": f"TC-{sheet}-{i}", "category": sheet,
                "test_steps": "1. 실행", "expected_result": "성공",
                "sheet_name": sheet, "priority": "보통",
            })
            assert rr.status_code == 201, rr.text

    yield pid
    requests.delete(f"{BASE}/api/projects/{pid}", headers=h)


def _export(token, pid, split=False):
    r = requests.get(f"{BASE}/api/projects/{pid}/testcases/export",
                     headers=auth(token), params={"split_sheets": str(split).lower()})
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def _tc_ids(ws):
    """헤더 아래 행에서 TC ID 칸만 뽑는다."""
    ids = []
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        values = [c.value for c in row]
        if header_row is None:
            if "TC ID" in values:
                header_row = row[0].row
                col = values.index("TC ID")
            continue
        if row[col].value:
            ids.append(row[col].value)
    return ids


def test_통합모드는_시트_차례대로_모은다(token, project):
    wb = _export(token, project, split=False)
    ws = wb.worksheets[0]
    got = _tc_ids(ws)

    assert got == [
        "TC-다시트-1", "TC-다시트-2",
        "TC-가시트-1", "TC-가시트-2",
        "TC-나시트-1", "TC-나시트-2",
    ], f"화면의 시트 차례를 따르지 않았다: {got}"


def test_시트분리모드의_탭_차례가_화면과_같다(token, project):
    wb = _export(token, project, split=True)
    assert wb.sheetnames == ["다시트", "가시트", "나시트"], (
        f"탭 차례가 화면과 다르다: {wb.sheetnames}"
    )
