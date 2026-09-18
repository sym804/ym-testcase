"""런 엑셀에 실리는 열

화면과 파일이 같은 것을 보여야 한다. 수행 화면에 사전조건 열이 없었고 런 엑셀에도
없었다(SYM-108). 엑셀 쪽은 headers 에서 빠진 것에 더해 load_only 에서도 빠져 있어
열을 헤더에만 더하면 행마다 지연 로딩이 붙는다. 그래서 값까지 확인한다.

실행: cd backend && python -m pytest tests_unit/test_run_export_columns.py -q
"""
import asyncio
import io
import os
import sys

import pytest

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from models import (
    Base, Project, TestCase, TestCaseSheet, TestResult, TestRun, TestResultValue, User,
)
from routes.reports import report_excel
from routes.testruns import export_testrun_excel

# 참조를 담은 사전조건. 줄바꿈은 chr(10) 으로 만든다.
PRECONDITION = chr(10).join(["1. 로그인한 상태", "2. TC-002 의 사전조건 참조"])


@pytest.fixture
def db(tmp_path):
    engine = create_engine(
        f"sqlite:///{(tmp_path / 'run_export.db').as_posix()}",
        connect_args={"check_same_thread": False},
    )

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_conn, _):
        cur = dbapi_conn.cursor()
        cur.execute("PRAGMA foreign_keys=ON")
        cur.close()

    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    yield s
    s.close()
    engine.dispose()


@pytest.fixture
def seeded(db):
    user = User(username="admin", password_hash="x", display_name="관리자", role="admin")
    db.add(user)
    db.flush()

    project = Project(name="P", created_by=user.id)
    db.add(project)
    db.flush()

    db.add(TestCaseSheet(project_id=project.id, name="결제", sort_order=0, is_folder=False))

    tc = TestCase(
        project_id=project.id, no=1, tc_id="TC-001", category="결제",
        priority="High", test_type="Android", precondition=PRECONDITION,
        test_steps="1. 결제한다", expected_result="성공",
        sheet_name="결제", created_by=user.id,
    )
    db.add(tc)
    db.flush()

    run = TestRun(project_id=project.id, name="결제 회귀", round=1, created_by=user.id)
    db.add(run)
    db.flush()

    db.add(TestResult(
        test_run_id=run.id, test_case_id=tc.id,
        result=TestResultValue.PASS, executed_by=user.id,
    ))
    db.commit()
    return project, run, user


def _read_body(resp):
    """StreamingResponse 는 동기 버퍼를 받아도 body_iterator 를 async 로 내준다."""
    async def collect():
        chunks = []
        async for c in resp.body_iterator:
            chunks.append(c if isinstance(c, bytes) else str(c).encode())
        return b"".join(chunks)

    return asyncio.run(collect())


def _sheet(db, seeded):
    project, run, user = seeded
    resp = export_testrun_excel(
        project_id=project.id, run_id=run.id, db=db, current_user=user,
    )
    return load_workbook(io.BytesIO(_read_body(resp))).active


def test_사전조건_열이_있다(db, seeded):
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    assert "Precondition" in headers, f"런 엑셀 헤더에 사전조건이 없다: {headers}"


def test_사전조건_값이_실린다(db, seeded):
    """헤더만 더하고 load_only 를 안 고치면 값이 비거나 지연 로딩이 붙는다."""
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    col = headers.index("Precondition") + 1
    assert ws.cell(row=2, column=col).value == PRECONDITION


def test_열_차례가_의도대로다(db, seeded):
    """이름으로만 찾는 검사는 순서가 뒤바뀌어도 통과한다. 차례를 못박는다.

    ★절대 인덱스(headers[7] == ...)로 적지 않는다. 열을 하나 더할 때마다 고쳐야
      해서, 고치다 보면 검사가 의도가 아니라 현재 구현을 베끼게 된다.
      실제로 Platform 을 넣을 때 그 형태였던 검사가 깨졌다.
    """
    ws = _sheet(db, seeded)
    h = [c.value for c in ws[1]]
    pos = {name: i for i, name in enumerate(h)}

    # TC 에서 오는 열은 TC 관리 화면 차례를 따른다
    assert pos["Priority"] < pos["Platform"] < pos["Precondition"] < pos["Test Steps"] < pos["Expected Result"]
    # 수행 결과는 TC 정보 뒤에 온다
    assert pos["Expected Result"] < pos["Result"] < pos["Actual Result"] < pos["Issue Link"]
    assert h[0] == "No"
    assert h[-1] == "Remarks"


def test_기존_열이_밀려나지_않는다(db, seeded):
    """열을 끼우면 뒤 열의 인덱스가 바뀐다. 값이 제 열에 들어가는지 본다."""
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
    assert row["TC ID"] == "TC-001"
    assert row["Priority"] == "High"
    assert row["Test Steps"] == "1. 결제한다"
    assert row["Expected Result"] == "성공"
    assert row["Result"] == "PASS"


def _build_run(db, tc_count):
    """TC 가 tc_count 개인 프로젝트와 런을 새로 만든다."""
    user = db.query(User).first()
    project = Project(name=f"P{tc_count}", created_by=user.id)
    db.add(project)
    db.flush()
    db.add(TestCaseSheet(project_id=project.id, name="결제", sort_order=0, is_folder=False))

    run = TestRun(project_id=project.id, name=f"런{tc_count}", round=1, created_by=user.id)
    db.add(run)
    db.flush()

    for i in range(tc_count):
        tc = TestCase(
            project_id=project.id, no=i + 1, tc_id=f"T-{tc_count}-{i}", category="결제",
            priority="High", precondition=PRECONDITION, test_steps="1. 실행",
            expected_result="성공", sheet_name="결제", created_by=user.id,
        )
        db.add(tc)
        db.flush()
        db.add(TestResult(
            test_run_id=run.id, test_case_id=tc.id,
            result=TestResultValue.PASS, executed_by=user.id,
        ))
    db.commit()
    return project, run, user


def _count_selects(db, tc_count, exporter=export_testrun_excel):
    project, run, user = _build_run(db, tc_count)
    seen = []

    engine = db.get_bind()

    def _rec(conn, cursor, statement, params, context, executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            seen.append(statement)

    event.listen(engine, "before_cursor_execute", _rec)
    try:
        db.expire_all()   # 캐시된 객체가 있으면 쿼리가 안 나가 비교가 무의미해진다
        _read_body(exporter(
            project_id=project.id, run_id=run.id, db=db, current_user=user,
        ))
    finally:
        event.remove(engine, "before_cursor_execute", _rec)
    return len(seen)


def test_행이_늘어도_조회_횟수가_같다(db, seeded):
    """load_only 에서 precondition 이 빠지면 행마다 지연 로딩이 한 번씩 붙는다.

    ★값만 보는 위 테스트들은 이것을 못 잡는다. 지연 로딩으로도 값은 나오기 때문이다.
      실제로 load_only 에서 precondition 을 빼고 돌려 보면 값 검증 4건이 모두 통과한다.
      그래서 행 수를 늘렸을 때 조회 횟수가 그대로인지로 본다.
    """
    few = _count_selects(db, 2)
    many = _count_selects(db, 6)
    assert few == many, (
        f"행이 2개일 때 {few}회, 6개일 때 {many}회 조회했다. "
        "load_only 에서 빠진 칸이 있어 행마다 지연 로딩이 붙는다"
    )


# 긴 글이 들어가는 열과 짧은 값만 들어가는 열. 폭이 밀리면 이 성질이 뒤집힌다.
WIDE_COLUMNS = {"Precondition", "Test Steps", "Expected Result", "Actual Result", "Remarks"}
NARROW_COLUMNS = {"No", "Type", "Result", "Duration(sec)", "Platform"}


def test_긴_열은_넓고_짧은_열은_좁다(db, seeded):
    """열을 끼우고 col_widths 를 안 고치면 폭이 한 칸씩 밀린다.

    값은 제 열에 들어가므로 다른 테스트가 못 잡는다. 실제로 이번에 Precondition 을
    넣고 col_widths 를 14개로 둔 채 넘어갔다가 QA 에서 잡혔다. 그때 Expected 가
    폭 10, Result 가 폭 30 을 받았다.

    ★"폭이 설정됐는지" 로는 못 잡는다. openpyxl 은 설정하지 않은 열에도 기본 폭
      13.0 을 돌려주기 때문이다. 실측으로 확인했고, 그래서 성질로 본다.
    """
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    width = {h: ws.column_dimensions[get_column_letter(i)].width for i, h in enumerate(headers, 1)}

    for h in sorted(WIDE_COLUMNS):
        assert width[h] >= 20, f"{h} 이 좁다({width[h]}). col_widths 가 밀렸다"
    for h in sorted(NARROW_COLUMNS):
        assert width[h] <= 15, f"{h} 이 넓다({width[h]}). col_widths 가 밀렸다"


def test_결과_색이_결과_열에_칠해진다(db, seeded):
    """열을 끼울 때 result_cell 의 열 번호를 안 고치면 엉뚱한 칸이 칠해진다."""
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    col = headers.index("Result") + 1
    cell = ws.cell(row=2, column=col)
    assert cell.value == "PASS"
    # 색상 상수는 두 파일이 다르다. 값을 박지 않고 "칠해졌는가" 로 본다.
    assert cell.fill.patternType == "solid", "결과 칸이 칠해지지 않았다"
    # 옆 칸까지 물들지 않았는지 본다
    assert ws.cell(row=2, column=col - 1).fill.start_color.rgb != cell.fill.start_color.rgb


def test_사전조건도_수식으로_읽히지_않는다(db):
    """새 열도 safe_cell 을 거쳐야 한다. 안 거치면 여는 쪽에서 수식이 실행된다(CWE-1236)."""
    user = User(username="u2", password_hash="x", display_name="u2", role="admin")
    db.add(user)
    db.flush()
    project = Project(name="P-inj", created_by=user.id)
    db.add(project)
    db.flush()
    db.add(TestCaseSheet(project_id=project.id, name="결제", sort_order=0, is_folder=False))
    tc = TestCase(
        project_id=project.id, no=1, tc_id="TC-INJ", category="결제", priority="High",
        precondition="=cmd|' /c calc'!A1", test_steps="1. 실행",
        expected_result="성공", sheet_name="결제", created_by=user.id,
    )
    db.add(tc)
    db.flush()
    run = TestRun(project_id=project.id, name="주입", round=1, created_by=user.id)
    db.add(run)
    db.flush()
    db.add(TestResult(
        test_run_id=run.id, test_case_id=tc.id,
        result=TestResultValue.PASS, executed_by=user.id,
    ))
    db.commit()

    ws = _sheet(db, (project, run, user))
    headers = [c.value for c in ws[1]]
    got = ws.cell(row=2, column=headers.index("Precondition") + 1).value
    assert not got.startswith("="), f"수식으로 읽힐 값이 그대로 남았다: {got!r}"
    assert "cmd" in got, "값 자체가 사라졌다. 지우는 것이 아니라 앞따옴표를 붙이는 것이다"


# ── 리포트 엑셀과의 계약 ────────────────────────────────────────────────────
# 같은 런을 두 파일로 뽑는다. 한쪽에만 열을 더하면 조용히 갈라진다.
# 수행 엑셀에 사전조건을 넣으면서 리포트를 빠뜨렸다가 QA 에서 잡혔다.

# 두 파일이 같은 것을 가리키면서 이름만 다른 열. 비교할 때 맞춰 준다.
HEADER_ALIASES = {"Steps": "Test Steps"}

# 리포트에만 없어도 되는 열. 리포트는 결과 요약이라 수행 메타는 빼고 본다.
RUN_ONLY = {"Duration(sec)"}


def _report_sheet(db, seeded):
    project, run, user = seeded
    resp = report_excel(
        project_id=project.id, run_id=run.id, db=db, current_user=user,
    )
    return load_workbook(io.BytesIO(_read_body(resp)))["Results"]


def test_리포트에도_사전조건이_있다(db, seeded):
    ws = _report_sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    assert "Precondition" in headers, f"리포트 Results 시트에 사전조건이 없다: {headers}"


def test_리포트의_사전조건_값이_실린다(db, seeded):
    ws = _report_sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("Precondition") + 1).value == PRECONDITION


def test_두_파일의_TC_열_집합이_같다(db, seeded):
    """열 이름 집합으로 묶는다. 한쪽에만 열을 더하면 여기서 걸린다."""
    run_headers = {c.value for c in _sheet(db, seeded)[1]}
    rep_headers = {HEADER_ALIASES.get(c.value, c.value) for c in _report_sheet(db, seeded)[1]}

    only_run = run_headers - rep_headers - RUN_ONLY
    only_rep = rep_headers - run_headers
    assert not only_run, f"수행 엑셀에만 있는 열: {sorted(only_run)}"
    assert not only_rep, f"리포트에만 있는 열: {sorted(only_rep)}"


def test_리포트도_결과_열에_색을_칠한다(db, seeded):
    """열 번호를 박아 두면 앞에 열을 끼울 때 엉뚱한 칸이 칠해진다."""
    ws = _report_sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    col = headers.index("Result") + 1
    cell = ws.cell(row=2, column=col)
    assert cell.value == "PASS"
    assert cell.fill.patternType == "solid", "결과 칸이 칠해지지 않았다"
    assert ws.cell(row=2, column=col - 1).fill.start_color.rgb != cell.fill.start_color.rgb


def test_런_엑셀에_Platform_이_있다(db, seeded):
    """사전조건과 같은 모양의 누락이었다(SYM-109)."""
    ws = _sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    assert "Platform" in headers, f"런 엑셀에 Platform 이 없다: {headers}"
    assert ws.cell(row=2, column=headers.index("Platform") + 1).value == "Android"


def test_리포트에도_Platform_이_있다(db, seeded):
    ws = _report_sheet(db, seeded)
    headers = [c.value for c in ws[1]]
    assert "Platform" in headers, f"리포트에 Platform 이 없다: {headers}"
    assert ws.cell(row=2, column=headers.index("Platform") + 1).value == "Android"


def test_리포트도_행이_늘어도_조회_횟수가_같다(db, seeded):
    """런 엑셀과 대칭으로 건다. 한쪽에만 걸어 두면 다른 쪽 load_only 가 새어도 모른다."""
    few = _count_selects(db, 2, report_excel)
    many = _count_selects(db, 6, report_excel)
    assert few == many, (
        f"행이 2개일 때 {few}회, 6개일 때 {many}회 조회했다. "
        "reports.py 의 load_only 에서 빠진 칸이 있다"
    )
