"""리포트(JSON/PDF/엑셀)에 실리는 내용

2026-09-23 실서버 팩트체크에서 나온 것들을 고정한다.
- BLOCK 만 있는 수행이 "실패 항목이 없습니다" 로 나왔다. 상세 목록이 FAIL 만 모았다.
- PDF 요약 표가 216mm 로 A4 본문(190mm)을 넘어 합격률 칸이 "Pas" 로 잘렸다.
- 엑셀 합격률이 문자열 "92.0%" 였다.
- 관련 이슈 목록이 set 을 거쳐 재시작마다 순서가 바뀔 수 있었다.

실행: cd backend && python -m pytest tests_unit/test_report_content.py -q
"""
import asyncio
import io
import os
import sys
from datetime import datetime, timedelta

import pytest

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from models import (
    Base, Project, TestCase, TestCaseSheet, TestResult, TestRun, TestResultValue, User,
)
from routes.reports import (
    _load_pdf_font, priority_sort_key, report_excel, report_filename, report_json, report_pdf,
)

R = TestResultValue


@pytest.fixture
def db(tmp_path):
    engine = create_engine(
        f"sqlite:///{(tmp_path / 'report.db').as_posix()}",
        connect_args={"check_same_thread": False},
    )

    @event.listens_for(engine, "connect")
    def _fk_on(dbapi_conn, _):
        cur = dbapi_conn.cursor()
        cur.execute("PRAGMA foreign_keys=ON")
        cur.close()

    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    yield s
    s.close()
    engine.dispose()


def _make(db, results, *, prev_results=None, project_name="P"):
    """results: [(tc_id, priority, category, 결과, 실제결과, 이슈링크, 수행자)]

    prev_results 를 주면 같은 TC 로 한 시간 앞선 런을 먼저 만든다.
    """
    admin = User(username="admin", password_hash="x", display_name="관리자", role="admin")
    tester = User(username="t1", password_hash="x", display_name="테스터", role="user")
    db.add_all([admin, tester])
    db.flush()
    users = {"admin": admin, "t1": tester}

    project = Project(name=project_name, created_by=admin.id)
    db.add(project)
    db.flush()
    db.add(TestCaseSheet(project_id=project.id, name="S", sort_order=0, is_folder=False))

    tcs = {}
    for no, (tc_id, priority, category, *_rest) in enumerate(results, 1):
        tc = TestCase(
            project_id=project.id, no=no, tc_id=tc_id, priority=priority, category=category,
            test_steps="1. 실행", expected_result="성공", sheet_name="S", created_by=admin.id,
        )
        db.add(tc)
        db.flush()
        tcs[tc_id] = tc

    base = datetime(2026, 9, 1, 10, 0)

    def add_run(name, when, rows):
        run = TestRun(project_id=project.id, name=name, round=1, created_by=admin.id, created_at=when)
        db.add(run)
        db.flush()
        for tc_id, result, actual, link, who in rows:
            db.add(TestResult(
                test_run_id=run.id, test_case_id=tcs[tc_id].id, result=result,
                actual_result=actual, issue_link=link, executed_by=users[who].id,
            ))
        return run

    if prev_results:
        add_run("이전", base, prev_results)
    run = add_run("이번", base + timedelta(hours=1), [(r[0], *r[3:]) for r in results])
    db.commit()
    return project, run, admin


def _json(db, made):
    project, run, user = made
    return report_json(project_id=project.id, run_id=run.id, db=db, current_user=user)


def _body(resp):
    async def collect():
        chunks = []
        async for c in resp.body_iterator:
            chunks.append(c if isinstance(c, bytes) else str(c).encode())
        return b"".join(chunks)

    return asyncio.run(collect())


# ── 상세 목록 ────────────────────────────────────────────────────────────────

def test_BLOCK_만_있어도_상세_목록에_실린다(db):
    made = _make(db, [
        ("A-1", "High", "검색", R.PASS, None, None, "t1"),
        ("A-2", "High", "검색", R.BLOCK, "입력창을 못 찾음", "PROJ-9", "t1"),
    ])
    data = _json(db, made)
    items = data["top_failures"]
    assert [i["test_case"]["tc_id"] for i in items] == ["A-2"]
    assert items[0]["result"] == "BLOCK", "결과를 FAIL 로 박으면 BLOCK 이 FAIL 로 보인다"
    assert items[0]["actual_result"] == "입력창을 못 찾음"
    assert data["jira_issues"] == ["PROJ-9"], "BLOCK 의 이슈 링크도 관련 이슈에 들어가야 한다"


def test_상세_목록은_우선순위_높은_것부터(db):
    made = _make(db, [
        ("L-1", "Low", "c", R.FAIL, None, "X-3", "t1"),
        ("H-1", "High", "c", R.BLOCK, None, "X-1", "t1"),
        ("N-1", None, "c", R.FAIL, None, "X-1", "t1"),
        ("M-1", "보통", "c", R.FAIL, None, "X-2", "t1"),
    ])
    data = _json(db, made)
    assert [i["test_case"]["tc_id"] for i in data["top_failures"]] == ["H-1", "M-1", "L-1", "N-1"]
    # 이슈 목록은 항목 차례를 따르고 중복을 한 번만 싣는다
    assert data["jira_issues"] == ["X-1", "X-2", "X-3"]


def test_우선순위_정렬_키():
    names = ["낮음", None, "사용자정의", "매우 높음", "Medium", "high", "  "]
    ordered = sorted(names, key=priority_sort_key)
    assert ordered[:4] == ["매우 높음", "high", "Medium", "낮음"]
    assert ordered[4] == "사용자정의", "모르는 값은 아는 값 뒤, 빈 값 앞"
    assert set(ordered[5:]) == {None, "  "}


# ── 집계 ────────────────────────────────────────────────────────────────────

def test_우선순위별_집계와_합격률(db):
    made = _make(db, [
        ("A", "High", "c", R.PASS, None, None, "t1"),
        ("B", "High", "c", R.FAIL, None, None, "t1"),
        ("C", "High", "c", R.NS, None, None, "admin"),
        ("D", None, "c", R.NA, None, None, "t1"),
    ])
    rows = {p["priority"]: p for p in _json(db, made)["priority_summary"]}
    high = rows["High"]
    assert (high["total"], high["pass"], high["fail"], high["not_started"]) == (3, 1, 1, 1)
    assert high["pass_rate"] == 50.0, "분모는 수행분(pass+fail+block)이다(SYM-57)"
    assert rows[None]["pass_rate"] is None, "수행한 것이 없으면 0% 가 아니라 없음이다"


def test_분류별_합격률(db):
    made = _make(db, [
        ("A", "High", "결제", R.PASS, None, None, "t1"),
        ("B", "High", "결제", R.BLOCK, None, None, "t1"),
        ("C", "High", "검색", R.NS, None, None, "admin"),
    ])
    rows = {c["category"]: c for c in _json(db, made)["category_summary"]}
    assert rows["결제"]["pass_rate"] == 50.0
    assert rows["검색"]["pass_rate"] is None


def test_수행자는_미수행_행을_세지_않는다(db):
    """결과 행은 런을 만들 때 NS 로 미리 생긴다. 세면 런을 만든 사람이 다 한 것처럼 나온다."""
    made = _make(db, [
        ("A", "High", "c", R.PASS, None, None, "t1"),
        ("B", "High", "c", R.FAIL, None, None, "t1"),
        ("C", "High", "c", R.NS, None, None, "admin"),
    ])
    assert _json(db, made)["executors"] == [{"name": "테스터", "count": 2}]


def test_지워진_TC_의_결과도_센다(db):
    """리포트는 그때의 수행 기록이다. 지금 상태를 보는 대시보드와 기준이 다른 것은 의도다.

    지워진 TC 를 거르는 필터가 집계 하나에만 붙어도 요약과 표의 합이 어긋난다.
    """
    made = _make(db, [
        ("A", "High", "결제", R.PASS, None, None, "t1"),
        ("B", "High", "결제", R.FAIL, "지운 TC", None, "t1"),
    ])
    tc = db.query(TestCase).filter(TestCase.tc_id == "B").one()
    tc.deleted_at = datetime(2026, 9, 2)
    db.commit()

    data = _json(db, made)
    assert data["summary"]["total"] == 2
    assert data["summary"]["fail"] == 1
    assert {c["category"]: c["total"] for c in data["category_summary"]} == {"결제": 2}
    assert {p["priority"]: p["total"] for p in data["priority_summary"]} == {"High": 2}
    assert [i["test_case"]["tc_id"] for i in data["top_failures"]] == ["B"]


# ── 직전 수행 대비 ───────────────────────────────────────────────────────────

def test_직전_수행이_없으면_비교가_없다(db):
    made = _make(db, [("A", "High", "c", R.PASS, None, None, "t1")])
    assert _json(db, made)["comparison"] is None


def test_직전_수행_대비_회귀와_해결(db):
    made = _make(
        db,
        [
            ("A", "High", "c", R.FAIL, None, None, "t1"),   # PASS -> FAIL 회귀
            ("B", "Low", "c", R.PASS, None, None, "t1"),    # FAIL -> PASS 해결
            ("C", "High", "c", R.BLOCK, None, None, "t1"),  # PASS -> BLOCK 변경만
            ("D", "High", "c", R.PASS, None, None, "t1"),   # 그대로
        ],
        prev_results=[
            ("A", R.PASS, None, None, "t1"),
            ("B", R.FAIL, None, None, "t1"),
            ("C", R.PASS, None, None, "t1"),
            ("D", R.PASS, None, None, "t1"),
        ],
    )
    comp = _json(db, made)["comparison"]
    assert comp["previous_run"]["name"] == "이전"
    assert comp["common"] == 4
    assert comp["changed"] == 3
    # 판정 기준은 수행 비교 화면(CompareView)과 같다
    assert [i["tc_id"] for i in comp["regressions"]] == ["A"]
    assert [i["tc_id"] for i in comp["fixed"]] == ["B"]


def test_직전_런의_미수행은_변경으로_세지_않는다(db):
    """직전 런이 진행 중이면 대부분 NS 다. NS -> PASS 를 변경으로 세면 잡음만 남는다."""
    made = _make(
        db,
        [
            ("A", "High", "c", R.PASS, None, None, "t1"),   # NS -> PASS: 세지 않는다
            ("B", "High", "c", R.FAIL, None, None, "t1"),   # PASS -> FAIL: 퇴보
            ("C", "High", "c", R.NS, None, None, "admin"),  # PASS -> NS: 세지 않는다
        ],
        prev_results=[
            ("A", R.NS, None, None, "admin"),
            ("B", R.PASS, None, None, "t1"),
            ("C", R.PASS, None, None, "t1"),
        ],
    )
    comp = _json(db, made)["comparison"]
    assert comp["common"] == 1, "양쪽 모두 수행한 TC 는 B 하나다"
    assert comp["changed"] == 1
    assert [i["tc_id"] for i in comp["regressions"]] == ["B"]


def test_분류가_NULL_과_빈_문자열이면_한_줄로_합친다(db):
    made = _make(db, [
        ("A", "High", None, R.PASS, None, None, "t1"),
        ("B", "High", "", R.FAIL, None, None, "t1"),
    ])
    rows = _json(db, made)["category_summary"]
    assert [(r["category"], r["total"]) for r in rows] == [("Uncategorized", 2)]


# ── 파일 ────────────────────────────────────────────────────────────────────

def test_파일명에서_Windows_금지_문자를_바꾼다(db):
    made = _make(db, [("A", "High", "c", R.PASS, None, None, "t1")], project_name="웹/앱: QA?")
    _, run, _ = made
    assert report_filename("웹/앱: QA?", run, "pdf") == "웹_앱_ QA__Report_R1.pdf"


def test_엑셀_합격률은_숫자다(db):
    made = _make(db, [
        ("A", "High", "c", R.PASS, None, None, "t1"),
        ("B", "High", "c", R.BLOCK, None, None, "t1"),
    ])
    project, run, user = made
    ws = load_workbook(io.BytesIO(_body(
        report_excel(project_id=project.id, run_id=run.id, db=db, current_user=user)
    )))["Summary"]
    rates = [c for row in ws.iter_rows() for c in row if c.number_format == "0.0%"]
    assert rates, "합격률 칸에 백분율 서식이 없다"
    assert all(isinstance(c.value, (int, float)) or c.value is None for c in rates)
    assert 0.5 in [c.value for c in rates], "전체 합격률 50% 가 0.5 로 들어가야 한다"
    titles = {c.value for row in ws.iter_rows() for c in row}
    assert {"Priority Breakdown", "Category Breakdown"} <= titles


def test_엑셀_요약의_분류명도_수식으로_읽히지_않는다(db):
    made = _make(db, [("A", "=1+1", "=HYPERLINK(1)", R.FAIL, None, None, "t1")])
    project, run, user = made
    ws = load_workbook(io.BytesIO(_body(
        report_excel(project_id=project.id, run_id=run.id, db=db, current_user=user)
    )))["Summary"]
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert not [v for v in values if v.startswith("=")], values


def test_PDF_는_레포에_넣은_한글_폰트를_쓴다():
    """시스템 폰트에 기대면 한글 폰트 없는 리눅스(CI 포함)에서 Helvetica 로 떨어지고,
    한글이 한 글자만 있어도 PDF 가 500 이 된다. 예전에는 이 테스트를 그 환경에서
    skip 했는데, skip 하면 바로 그 환경의 결함을 못 본다."""
    from fpdf import FPDF

    fonts = os.path.join(BACKEND, "fonts")
    for name in ("Pretendard-Regular.ttf", "Pretendard-Bold.ttf", "OFL.txt"):
        assert os.path.exists(os.path.join(fonts, name)), f"{name} 이 backend/fonts 에 없다"
    pdf = FPDF()
    assert _load_pdf_font(pdf) == "KoreanFont"
    assert "pretendard" in str(pdf.fonts["koreanfont"].ttffile).lower(), "레포 폰트보다 시스템 폰트가 먼저 잡혔다"


def test_PDF_합격률_칸이_잘리지_않는다(db):
    from pypdf import PdfReader

    rows = [(f"T-{i}", "High", "c", R.PASS, None, None, "t1") for i in range(23)]
    rows += [(f"B-{i}", "High", "c", R.BLOCK, "막힘 사유", None, "t1") for i in range(2)]
    made = _make(db, rows)
    project, run, user = made
    reader = PdfReader(io.BytesIO(_body(
        report_pdf(project_id=project.id, run_id=run.id, db=db, current_user=user)
    )))
    # ★글자가 있는지만 보면 안 된다. 페이지 밖으로 밀린 글자도 텍스트 추출에는 나온다
    #   (옛 216mm 배치에서도 "92.0%" in text 는 참이었다). 시작 x 좌표로 본다.
    #   옛 배치의 합격률 값은 x=590pt 에서 시작했다(A4 폭 595pt, 오른쪽 여백선 567pt).
    xs = []

    def visit(text, cm, tm, font_dict, font_size):
        if text.strip() == "92.0%":
            xs.append(tm[4] * cm[0] + cm[4])

    page = reader.pages[0]
    text = page.extract_text(visitor_text=visit)
    assert "PASS Rate" in text
    right_margin = float(page.mediabox.width) - 10 / 25.4 * 72
    assert xs, "합격률 값이 PDF 에 없다"
    assert xs[0] + 20 <= right_margin, f"합격률 값이 x={xs[0]:.0f}pt 에서 시작해 여백선({right_margin:.0f}pt)을 넘는다"
    assert "B-0" in text and "[BLOCK]" in text, "BLOCK 상세가 PDF 에 없다"
    fonts = {
        str(f["/BaseFont"])
        for p in reader.pages
        for f in (p["/Resources"]["/Font"].get_object().values())
    }
    assert len(fonts) >= 2, f"굵은체가 따로 임베드되지 않았다: {fonts}"
