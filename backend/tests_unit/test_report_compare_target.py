"""리포트의 비교 대상

- 자동: 같은 이름의 이전 회차(R2 면 R1). 없으면 비교하지 않는다.
- 선택: 수행에 저장한 compare_run_id 가 있으면 그것. PDF·엑셀도 같은 대상과 비교한다.

예전에는 바로 전에 만든 수행을 이름과 상관없이 골라, 계정·세션 정책 R1 이 관계없는
Full 테스트와 비교되어 겹치는 TC 0건인 결과가 실렸다(2026-09-30 실측).

실행: cd backend && python -m pytest tests_unit/test_report_compare_target.py -q
"""
import asyncio
import io
import os
import sys
from datetime import datetime, timedelta

import pytest

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BACKEND not in sys.path:
    sys.path.insert(0, BACKEND)

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from models import Project, TestCase, TestCaseSheet, TestResult, TestResultValue, TestRun, User
from routes.reports import report_excel, report_json, report_pdf
from routes.testruns import update_testrun
from schemas import TestRunUpdate

R = TestResultValue
BASE = datetime(2026, 9, 1, 10, 0)


@pytest.fixture
def db(pg_engine, tmp_path):
    engine = pg_engine

    s = sessionmaker(bind=engine)()
    yield s
    s.close()
    engine.dispose()


@pytest.fixture
def env(db):
    user = User(username="a", password_hash="x", display_name="관리자", role="admin")
    db.add(user)
    db.flush()
    project = Project(name="P", created_by=user.id)
    db.add(project)
    db.flush()
    db.add(TestCaseSheet(project_id=project.id, name="S", sort_order=0, is_folder=False))
    tc = TestCase(project_id=project.id, no=1, tc_id="A-1", priority="High", category="c",
                  test_steps="1", expected_result="2", sheet_name="S", created_by=user.id)
    db.add(tc)
    db.commit()
    return db, project, user, tc


def _run(env, name, round_, hours, result=R.PASS):
    db, project, user, tc = env
    run = TestRun(project_id=project.id, name=name, round=round_, created_by=user.id,
                  created_at=BASE + timedelta(hours=hours))
    db.add(run)
    db.flush()
    db.add(TestResult(test_run_id=run.id, test_case_id=tc.id, result=result, executed_by=user.id))
    db.commit()
    return run


def _comp(env, run):
    db, project, user, _ = env
    return report_json(project_id=project.id, run_id=run.id, db=db, current_user=user)["comparison"]


def _set(env, run, target_id):
    db, project, user, _ = env
    return update_testrun(project.id, run.id, TestRunUpdate(compare_run_id=target_id), db=db, current_user=user)


def _body(resp):
    async def collect():
        return b"".join([c async for c in resp.body_iterator])
    return asyncio.run(collect())


def test_R2_는_같은_이름의_R1_과_자동_비교한다(env):
    r1 = _run(env, "계정·세션 정책", 1, 0, R.PASS)
    _run(env, "Full 테스트", 1, 1)  # R1 과 R2 사이에 만든 다른 수행
    r2 = _run(env, "계정·세션 정책", 2, 2, R.FAIL)
    comp = _comp(env, r2)
    assert comp["previous_run"]["id"] == r1.id
    assert comp["mode"] == "auto"
    assert [i["tc_id"] for i in comp["regressions"]] == ["A-1"]


def test_R3_은_가장_가까운_이전_회차와_비교한다(env):
    _run(env, "수행", 1, 0)
    r2 = _run(env, "수행", 2, 1)
    r3 = _run(env, "수행", 3, 2)
    assert _comp(env, r3)["previous_run"]["id"] == r2.id


def test_R1_은_자동으로_비교하지_않는다(env):
    """이름이 다른 수행과 비교하면 겹치는 TC 0건인 결과가 실린다."""
    _run(env, "Full 테스트", 1, 0)
    r1 = _run(env, "계정·세션 정책", 1, 1)
    assert _comp(env, r1) is None


def test_선택한_대상이_자동보다_먼저다(env):
    full = _run(env, "Full 테스트", 1, 0, R.FAIL)
    _run(env, "정책", 1, 1)
    r2 = _run(env, "정책", 2, 2, R.PASS)
    _set(env, r2, full.id)
    comp = _comp(env, r2)
    assert (comp["previous_run"]["id"], comp["mode"]) == (full.id, "manual")
    assert [i["tc_id"] for i in comp["fixed"]] == ["A-1"]


def test_null_로_자동으로_되돌린다(env):
    r1 = _run(env, "정책", 1, 0)
    other = _run(env, "Full", 1, 1)
    r2 = _run(env, "정책", 2, 2)
    _set(env, r2, other.id)
    assert _set(env, r2, None).compare_run_id is None
    assert _comp(env, r2)["previous_run"]["id"] == r1.id


def test_자기_자신과_다른_프로젝트_수행은_거절한다(env):
    db, project, user, _ = env
    run = _run(env, "정책", 1, 0)
    with pytest.raises(HTTPException) as e:
        _set(env, run, run.id)
    assert e.value.status_code == 400
    other_project = Project(name="Q", created_by=user.id)
    db.add(other_project)
    db.flush()
    foreign = TestRun(project_id=other_project.id, name="x", round=1, created_by=user.id)
    db.add(foreign)
    db.commit()
    with pytest.raises(HTTPException) as e:
        _set(env, run, foreign.id)
    assert e.value.status_code == 400


def test_비교_대상을_지우면_자동으로_돌아간다(env):
    db, *_ = env
    r1 = _run(env, "정책", 1, 0)
    other = _run(env, "Full", 1, 1)
    r2 = _run(env, "정책", 2, 2)
    _set(env, r2, other.id)
    db.delete(other)
    db.commit()
    db.refresh(r2)
    assert r2.compare_run_id is None, "ON DELETE SET NULL"
    assert _comp(env, r2)["previous_run"]["id"] == r1.id


def test_PDF_와_엑셀도_선택한_대상과_비교한다(env):
    db, project, user, _ = env
    full = _run(env, "Full 테스트", 1, 0, R.FAIL)
    r1 = _run(env, "정책", 1, 1, R.PASS)
    _set(env, r1, full.id)

    from pypdf import PdfReader
    text = "".join(pg.extract_text() for pg in PdfReader(io.BytesIO(_body(
        report_pdf(project_id=project.id, run_id=r1.id, db=db, current_user=user)))).pages)
    assert "비교 대상: Full 테스트 (R1)" in text

    ws = load_workbook(io.BytesIO(_body(
        report_excel(project_id=project.id, run_id=r1.id, db=db, current_user=user))))["Summary"]
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "비교 대상: Full 테스트 (R1)" in values


def test_이전_FAIL_이_어떻게_됐는지_세고_FAIL_그대로인_TC_도_상세에_싣는다(env):
    """FAIL 8건 중 7건 고침이면 남은 1건(FAIL 그대로)도 보여야 한다. 변경만 세면 빠진다."""
    db, project, user, tc_a = env
    tc_a.category = "인증"

    def tc(no, tc_id, priority, category):
        t = TestCase(project_id=project.id, no=no, tc_id=tc_id, priority=priority, category=category,
                     test_steps="1", expected_result="2", sheet_name="S", created_by=user.id)
        db.add(t)
        db.flush()
        return t

    tc_b, tc_c, tc_d = tc(2, "A-2", "Low", "세션"), tc(3, "A-3", "High", "인증"), tc(4, "A-4", "High", "인증")
    # _run 은 tc_a 결과 한 줄을 넣는다(R1 FAIL, R2 PASS = 개선)
    r1 = _run(env, "정책", 1, 0, R.FAIL)
    r2 = _run(env, "정책", 2, 1, R.PASS)
    for t, before, after in ((tc_b, R.FAIL, R.FAIL), (tc_c, R.FAIL, R.BLOCK), (tc_d, R.PASS, R.FAIL)):
        db.add(TestResult(test_run_id=r1.id, test_case_id=t.id, result=before, executed_by=user.id))
        db.add(TestResult(test_run_id=r2.id, test_case_id=t.id, result=after, executed_by=user.id))
    db.commit()
    from routes.run_issues import create_run_issue
    from schemas import RunIssueCreate
    create_run_issue(project.id, r2.id, RunIssueCreate(title="t", url="https://linear.app/s/issue/SF-1", tc_ids=["A-2"]),
                     db=db, current_user=user)

    comp = report_json(project_id=project.id, run_id=r2.id, db=db, current_user=user)["comparison"]
    assert comp["prev_fail"] == {"total": 3, "fixed": 1, "still": 1, "other": 1}
    assert comp["changed"] == 3 and len(comp["regressions"]) == 1 and len(comp["fixed"]) == 1
    # 차례: 퇴보 -> 미수정 -> 개선 -> 그 밖. 미수정(FAIL 그대로)은 변경이 아니지만 상세에 실린다
    assert [(c["tc_id"], c["kind"], c["before"], c["after"]) for c in comp["changes"]] == [
        ("A-4", "regression", "PASS", "FAIL"), ("A-2", "still", "FAIL", "FAIL"),
        ("A-1", "fixed", "FAIL", "PASS"), ("A-3", "other", "FAIL", "BLOCK"),
    ]
    assert comp["changes"][1]["issue_keys"] == ["SF-1"] and comp["changes"][1]["category"] == "세션"

    ws = load_workbook(io.BytesIO(_body(
        report_excel(project_id=project.id, run_id=r2.id, db=db, current_user=user))))["Summary"]
    values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    # 카드 넷이 표 머리글로 실린다. 다른 결과로 변경(FAIL -> BLOCK)은 카드 대신 변경 상세에서 본다
    head = next(r for r in ws.iter_rows() if r[1].value == "양쪽 모두 수행")
    assert [c.value for c in head[1:6]] == ["양쪽 모두 수행", "이전 실패 TC", "수정된 TC", "미수정 TC", "퇴보 TC"]
    assert [c.value for c in ws[head[0].row + 1][1:6]] == [4, 3, 1, 1, 1]
    assert "변경 상세" in values and "FAIL -> FAIL" in values and "미수정" in values

    from pypdf import PdfReader
    text = "".join(pg.extract_text() for pg in PdfReader(io.BytesIO(_body(
        report_pdf(project_id=project.id, run_id=r2.id, db=db, current_user=user)))).pages)
    assert "변경 상세 4" in text and "FAIL -> FAIL" in text and "SF-1" in text
