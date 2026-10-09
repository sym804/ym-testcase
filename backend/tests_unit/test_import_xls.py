"""TC 가져오기의 Excel 형식 판정.

openpyxl 은 Excel 97-2003 형식(OLE2/BIFF)을 읽지 못한다. 예전에는 읽기에서 'Failed to read Excel file'
로만 떨어졌다. 판정은 이름이 아니라 내용으로 한다. 이름만 .xls 이고 내용은 xlsx 인 파일은 다른 시스템의
내보내기에서 흔하고, openpyxl 이 그대로 읽는다.
"""
import io

import openpyxl
import pytest
from fastapi import HTTPException
from starlette.datastructures import UploadFile

from services.import_service import XLS_NOT_SUPPORTED, _load_workbook_from_upload

# .xls 파일은 OLE2 컨테이너 시그니처로 시작한다
OLE2_XLS = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "TC ID"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(name: str, data: bytes) -> UploadFile:
    return UploadFile(file=io.BytesIO(data), filename=name)


def test_xls_는_이유와_함께_거절한다():
    with pytest.raises(HTTPException) as e:
        _load_workbook_from_upload(_upload("old.xls", OLE2_XLS))
    assert e.value.status_code == 400
    assert e.value.detail == XLS_NOT_SUPPORTED


def test_이름만_xls_이고_내용이_xlsx_면_읽는다():
    wb = _load_workbook_from_upload(_upload("export.xls", _xlsx_bytes()))
    assert wb.active["A1"].value == "TC ID"


def test_이름은_xlsx_인데_내용이_옛_형식이면_같은_이유로_거절한다():
    with pytest.raises(HTTPException) as e:
        _load_workbook_from_upload(_upload("renamed.xlsx", OLE2_XLS))
    assert e.value.detail == XLS_NOT_SUPPORTED


def test_확장자가_대문자인_xlsx_도_읽는다():
    wb = _load_workbook_from_upload(_upload("TC.XLSX", _xlsx_bytes()))
    assert wb.active["A1"].value == "TC ID"


def test_xlsx_가_아닌_확장자는_거절한다():
    with pytest.raises(HTTPException) as e:
        _load_workbook_from_upload(_upload("notes.txt", b"hello"))
    assert e.value.status_code == 400
