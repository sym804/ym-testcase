"""엑셀로 내보내는 값은 수식으로 실행되지 않는다 (CWE-1236).

TC 목록 내보내기에만 새니타이즈가 있었다. 수행 엑셀과 리포트 엑셀은 사용자가 쓴
값을 그대로 넣어, 파일을 여는 사람의 PC 에서 =HYPERLINK 나 DDE 가 실행될 수 있다.

실행: cd backend && TEST_PORT=8009 TEST_BASE_URL=http://127.0.0.1:8009 python -m pytest test_export_formula_injection.py -v
"""
import io
import os

import pytest
import requests
from openpyxl import load_workbook

BASE = os.getenv("TEST_BASE_URL", "http://127.0.0.1:8008")
ADMIN_PW = os.getenv("TEST_ADMIN_PASSWORD", "test1234")

import dev_db_guard

if dev_db_guard.DEV_DB_AT_RISK:
    pytest.skip(dev_db_guard.SKIP_REASON, allow_module_level=True)


#: 엑셀이 수식으로 읽는 시작 문자. 앞따옴표를 붙여 텍스트로 만들어야 한다.
DANGEROUS = "=HYPERLINK(\"http://evil\",\"click\")"


def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": ADMIN_PW})
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture
def run_with_formula(token):
    """수식으로 읽힐 값이 TC 와 결과에 들어 있는 수행."""
    h = auth(token)
    r = requests.post(f"{BASE}/api/projects", headers=h, json={"name": "__formula__"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]

    rs = requests.post(f"{BASE}/api/projects/{pid}/testcases/sheets", headers=h,
                       json={"name": "보안", "parent_id": None, "is_folder": False})
    assert rs.status_code in (200, 201), rs.text

    rt = requests.post(f"{BASE}/api/projects/{pid}/testcases", headers=h, json={
        "no": 1, "tc_id": "TC-SEC-001", "category": DANGEROUS,
        "test_steps": DANGEROUS, "expected_result": "성공",
        "sheet_name": "보안", "priority": "High",
    })
    assert rt.status_code == 201, rt.text
    tc_id = rt.json()["id"]

    rr = requests.post(f"{BASE}/api/projects/{pid}/testruns", headers=h, json={"name": "sec run"})
    assert rr.status_code == 201, rr.text
    run_id = rr.json()["id"]

    rres = requests.post(f"{BASE}/api/projects/{pid}/testruns/{run_id}/results", headers=h, json=[{
        "test_case_id": tc_id, "result": "FAIL", "actual_result": DANGEROUS, "remarks": DANGEROUS,
    }])
    assert rres.status_code in (200, 201), rres.text

    yield pid, run_id
    requests.delete(f"{BASE}/api/projects/{pid}", headers=h)


def _cells(content, sheet=None):
    wb = load_workbook(io.BytesIO(content))
    ws = wb[sheet] if sheet else wb.active
    out = []
    for row in ws.iter_rows(values_only=True):
        out.extend(v for v in row if isinstance(v, str))
    return out


def test_lead_characters_are_all_escaped():
    """막아야 하는 시작 문자를 하나도 빠뜨리지 않았는지 본다.

    내보내기 경로를 다 돌지 않고 함수만 본다. 경로별 적용 여부는 아래 테스트가 본다.
    """
    from services.excel_safe import safe_cell

    missed = [c for c in ("=", "+", "-", "@", "\t", "\r", "\n") if not safe_cell(c + "x").startswith("'")]
    assert not missed, f"막지 못하는 시작 문자: {missed!r}"
    assert safe_cell("정상 값") == "정상 값", "위험하지 않은 값은 건드리지 않는다"
    assert safe_cell(None) is None
    assert safe_cell(3) == 3


def _assert_no_formula(cells):
    live = [v for v in cells if v.startswith(("=", "+", "-", "@", "\t", "\r", "\n"))]
    assert not live, f"수식으로 읽힐 셀이 남았다: {live[:3]}"
    # 값 자체는 남아 있어야 한다. 지우는 것이 아니라 앞따옴표를 붙이는 것이다.
    assert any(DANGEROUS in v for v in cells), "위험 문자열이 아예 사라졌다"


def test_run_export_escapes_formula(token, run_with_formula):
    pid, run_id = run_with_formula
    r = requests.get(f"{BASE}/api/projects/{pid}/testruns/{run_id}/export", headers=auth(token))
    assert r.status_code == 200, r.text

    _assert_no_formula(_cells(r.content))


def test_report_excel_escapes_formula(token, run_with_formula):
    pid, run_id = run_with_formula
    r = requests.get(f"{BASE}/api/projects/{pid}/reports/excel",
                     headers=auth(token), params={"run_id": run_id})
    assert r.status_code == 200, r.text

    _assert_no_formula(_cells(r.content, "Results"))


def test_testcase_export_still_escapes(token, run_with_formula):
    """이미 막혀 있던 TC 목록 내보내기가 그대로인지 함께 본다."""
    pid, _ = run_with_formula
    r = requests.get(f"{BASE}/api/projects/{pid}/testcases/export", headers=auth(token))
    assert r.status_code == 200, r.text

    _assert_no_formula(_cells(r.content))
